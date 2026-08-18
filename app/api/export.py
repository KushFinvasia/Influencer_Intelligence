"""CSV and Excel export endpoints with executive-ready formatting."""

from __future__ import annotations

import re
from datetime import datetime, timezone
from io import BytesIO

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy.ext.asyncio import AsyncSession

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.database.session import get_db
from app.schemas.search import SearchFilters
from app.services.creator_service import CreatorService
from app.utils.helpers import creators_to_csv

router = APIRouter(prefix="/api/export", tags=["Export"])


class DirectExportRequest(BaseModel):
    creators: list[dict]
    visible_columns: list[str] | None = None


COLUMN_DEFINITIONS = {
    "platform": {"label": "Platform", "align": "center", "width": 15},
    "name": {"label": "Creator Name", "align": "left", "width": 25},
    "followers": {"label": "Followers", "align": "right", "width": 16, "number_format": "#,##0"},
    "tier": {"label": "Audience Tier", "align": "center", "width": 18},
    "format": {"label": "Content Format", "align": "center", "width": 16},
    "avg_views": {"label": "Avg Views", "align": "right", "width": 16, "number_format": "#,##0"},
    "avg_likes": {"label": "Avg Likes", "align": "right", "width": 16, "number_format": "#,##0"},
    "avg_comments": {"label": "Avg Comments", "align": "right", "width": 16, "number_format": "#,##0"},
    "performance": {"label": "Performance Breakdown", "align": "left", "width": 38},
    "engagement": {"label": "Engagement Rate", "align": "right", "width": 16, "number_format": "0.00%"},
    "email": {"label": "Email", "align": "left", "width": 28},
    "phone": {"label": "Phone", "align": "left", "width": 18},
    "category": {"label": "Primary Category", "align": "left", "width": 24},
    "language": {"label": "Language", "align": "left", "width": 16},
    "social_handles": {"label": "Social Handles", "align": "left", "width": 22},
    "broker": {"label": "Broker Association", "align": "left", "width": 20},
    "website": {"label": "Website", "align": "left", "width": 25},
    "profile": {"label": "Profile Link", "align": "left", "width": 35},
    "updated": {"label": "Metrics Updated", "align": "center", "width": 20},
}


def _parse_numeric(val: str | int | float | None) -> int | float | str:
    """Parse string number representations (including 15M, 250K, 4.5M) to true numeric values for Excel."""
    if val is None or val == "-":
        return ""
    if isinstance(val, (int, float)):
        return val

    val_str = str(val).strip()
    # Check for percentage
    if val_str.endswith("%"):
        try:
            return float(val_str.rstrip("%")) / 100.0
        except ValueError:
            return val_str

    val_upper = val_str.upper()
    multiplier = 1
    if val_upper.endswith("K"):
        multiplier = 1_000
        val_str = val_str[:-1].strip()
    elif val_upper.endswith("M"):
        multiplier = 1_000_000
        val_str = val_str[:-1].strip()
    elif val_upper.endswith("B"):
        multiplier = 1_000_000_000
        val_str = val_str[:-1].strip()

    # Clean numbers like '1,250,000' or '1250000'
    cleaned = re.sub(r"[^\d.]", "", val_str)
    if cleaned:
        try:
            if "." in cleaned:
                num = float(cleaned) * multiplier
                return int(num) if num.is_integer() else round(num, 2)
            return int(int(cleaned) * multiplier)
        except ValueError:
            pass
    return val_str


def _extract_cell_value(creator: dict, col_id: str) -> tuple[any, str | None]:
    """Extract value and optional number format for a column."""
    if col_id == "platform":
        return creator.get("platform", "").upper(), None
    elif col_id == "name":
        return creator.get("name", "") or creator.get("handle", ""), None
    elif col_id == "followers":
        raw = creator.get("followers_raw") if creator.get("followers_raw") is not None else (creator.get("followers") or creator.get("follower_count"))
        return _parse_numeric(raw), "#,##0"
    elif col_id == "tier":
        return creator.get("bucket") or creator.get("tier") or creator.get("audience_bucket", ""), None
    elif col_id == "format":
        fmt = creator.get("content_format") or creator.get("format") or ""
        return fmt.capitalize(), None
    elif col_id == "avg_views":
        raw = creator.get("avg_views_raw") if creator.get("avg_views_raw") is not None else creator.get("avg_views")
        return _parse_numeric(raw), "#,##0"
    elif col_id == "avg_likes":
        raw = creator.get("avg_likes_raw") if creator.get("avg_likes_raw") is not None else creator.get("avg_likes")
        return _parse_numeric(raw), "#,##0"
    elif col_id == "avg_comments":
        raw = creator.get("avg_comments_raw") if creator.get("avg_comments_raw") is not None else creator.get("avg_comments")
        return _parse_numeric(raw), "#,##0"
    elif col_id == "performance":
        perf = creator.get("performance")
        if isinstance(perf, str) and perf != "-":
            return perf, None
        views = creator.get("avg_views", "-")
        likes = creator.get("avg_likes", "-")
        comments = creator.get("avg_comments", "-")
        parts = []
        if views != "-": parts.append(f"Views: {views}")
        if likes != "-": parts.append(f"Likes: {likes}")
        if comments != "-": parts.append(f"Comments: {comments}")
        return " | ".join(parts) if parts else "-", None
    elif col_id == "engagement":
        raw = creator.get("engagement_rate") or creator.get("engagement")
        num = _parse_numeric(raw)
        return num, "0.00%" if isinstance(num, (int, float)) else None
    elif col_id == "email":
        e = creator.get("email", "")
        return "" if e == "-" else e, None
    elif col_id == "phone":
        p = creator.get("phone", "")
        return "" if p == "-" else p, None
    elif col_id == "category":
        c = creator.get("category") or creator.get("primary_category", "")
        return "" if c == "-" else c, None
    elif col_id == "language":
        l = creator.get("language") or creator.get("primary_language", "")
        return "" if l == "-" else l, None
    elif col_id == "broker":
        b = creator.get("broker", "")
        return "" if b == "-" else b, None
    elif col_id == "social_handles":
        handles = creator.get("social_handles") or creator.get("handles")
        if isinstance(handles, list):
            return ", ".join(handles), None
        return handles or "", None
    elif col_id == "website":
        return creator.get("website", ""), None
    elif col_id == "profile":
        return creator.get("profile_url", ""), None
    elif col_id == "updated":
        return creator.get("metrics_calculated_at", ""), None
    else:
        return creator.get(col_id, ""), None


def build_styled_excel(creators: list[dict], visible_cols: list[str] | None = None) -> BytesIO:
    """Generate an executive-ready, formatted Excel workbook."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Creator Directory"

    # Determine columns to display
    if not visible_cols:
        visible_cols = [
            "platform", "name", "followers", "tier", "format",
            "avg_views", "avg_likes", "avg_comments", "engagement",
            "email", "phone", "category", "language", "broker", "profile"
        ]

    active_cols = [c for c in visible_cols if c in COLUMN_DEFINITIONS]
    if not active_cols:
        active_cols = list(COLUMN_DEFINITIONS.keys())

    # --- Styles ---
    # Fonts
    title_font = Font(name="Segoe UI", size=14, bold=True, color="0F172A")
    subtitle_font = Font(name="Segoe UI", size=9, italic=True, color="64748B")
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=9.5, color="1E293B")
    bold_data_font = Font(name="Segoe UI", size=9.5, bold=True, color="0F172A")
    link_font = Font(name="Segoe UI", size=9.5, color="2563EB", underline="single")
    bold_link_font = Font(name="Segoe UI", size=9.5, bold=True, color="2563EB", underline="single")

    # Fills
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    even_row_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    odd_row_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    # Borders
    thin_border_side = Side(border_style="thin", color="E2E8F0")
    header_border_side = Side(border_style="medium", color="334155")
    data_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    header_border = Border(left=thin_border_side, right=thin_border_side, top=header_border_side, bottom=header_border_side)

    # --- Row 1: Title Banner ---
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(active_cols))
    title_cell = ws.cell(row=1, column=1, value="Financial Creator Intelligence Directory")
    title_cell.font = title_font
    ws.row_dimensions[1].height = 24

    # --- Row 2: Subtitle / Export Info ---
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(active_cols))
    now_str = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    subtitle_cell = ws.cell(row=2, column=1, value=f"Generated on {now_str}  |  Total Filtered Creators: {len(creators)}")
    subtitle_cell.font = subtitle_font
    ws.row_dimensions[2].height = 18

    # --- Row 4: Table Headers ---
    header_row_idx = 4
    ws.row_dimensions[header_row_idx].height = 28

    for col_idx, col_id in enumerate(active_cols, start=1):
        col_def = COLUMN_DEFINITIONS[col_id]
        cell = ws.cell(row=header_row_idx, column=col_idx, value=col_def["label"])
        cell.font = header_font
        cell.fill = header_fill
        cell.border = header_border
        cell.alignment = Alignment(
            horizontal=col_def["align"],
            vertical="center",
            wrap_text=True,
        )

    # --- Data Rows ---
    start_data_row = 5
    for row_offset, creator in enumerate(creators):
        current_row = start_data_row + row_offset
        ws.row_dimensions[current_row].height = 22
        row_fill = even_row_fill if row_offset % 2 == 0 else odd_row_fill

        profile_link = (
            creator.get("profile_url")
            or creator.get("profile")
            or creator.get("url")
            or creator.get("channel_url")
        )

        for col_idx, col_id in enumerate(active_cols, start=1):
            col_def = COLUMN_DEFINITIONS[col_id]
            val, custom_fmt = _extract_cell_value(creator, col_id)

            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = data_font
            cell.fill = row_fill
            cell.border = data_border
            cell.alignment = Alignment(
                horizontal=col_def["align"],
                vertical="center",
                wrap_text=(col_id in ["performance", "social_handles", "category"]),
            )

            # Apply font styling for special columns
            if col_id == "name":
                if profile_link and isinstance(profile_link, str) and profile_link.startswith("http"):
                    cell.font = bold_link_font
                    cell.hyperlink = profile_link
                else:
                    cell.font = bold_data_font
            elif col_id in ["profile", "website"] and isinstance(val, str) and val.startswith("http"):
                cell.font = link_font
                cell.hyperlink = val

            # Apply number format if numeric
            fmt = custom_fmt or col_def.get("number_format")
            if fmt and isinstance(val, (int, float)):
                cell.number_format = fmt

    # --- Auto Column Widths ---
    for col_idx, col_id in enumerate(active_cols, start=1):
        col_letter = get_column_letter(col_idx)
        def_width = COLUMN_DEFINITIONS[col_id]["width"]

        # Calculate max content length
        max_len = len(COLUMN_DEFINITIONS[col_id]["label"])
        for row in range(start_data_row, start_data_row + min(len(creators), 100)):
            cell_val = str(ws.cell(row=row, column=col_idx).value or "")
            max_len = max(max_len, min(len(cell_val), 40))

        adjusted_width = max(def_width, max_len + 4)
        ws.column_dimensions[col_letter].width = adjusted_width

    # --- Freeze Panes & AutoFilter ---
    ws.freeze_panes = f"A{start_data_row}"
    last_col_letter = get_column_letter(len(active_cols))
    last_row = start_data_row + max(len(creators) - 1, 0)
    ws.auto_filter.ref = f"A{header_row_idx}:{last_col_letter}{last_row}"

    # Write to memory buffer
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


@router.post("/excel")
async def export_excel_post(request: DirectExportRequest):
    """Export frontend-filtered creators as a beautifully styled Excel workbook."""
    output = build_styled_excel(request.creators, request.visible_columns)
    filename = f"creators_export_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/csv")
async def export_csv(
    q: str | None = None,
    category: str | None = None,
    language: str | None = None,
    creator_type: str | None = None,
    broker: str | None = None,
    audience_bucket: int | None = None,
    platform: str | None = None,
    has_email: bool | None = None,
    status: str | None = None,
    min_score: float | None = None,
    db: AsyncSession = Depends(get_db),
):
    """Export filtered creators as a CSV file."""
    filters = SearchFilters(
        q=q,
        category=category,
        language=language,
        creator_type=creator_type,
        broker=broker,
        audience_bucket=audience_bucket,
        platform=platform,
        has_email=has_email,
        status=status,
        min_score=min_score,
        page=1,
        page_size=5000,
    )

    service = CreatorService(db)
    result = await service.search_creators(filters)

    rows = [r.model_dump() if hasattr(r, "model_dump") else r for r in result.results]
    csv_data = creators_to_csv(rows)

    return StreamingResponse(
        iter([csv_data]),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=creators_export.csv"},
    )


@router.get("/excel")
async def export_excel(
    q: str | None = None,
    category: str | None = None,
    language: str | None = None,
    creator_type: str | None = None,
    broker: str | None = None,
    audience_bucket: int | None = None,
    platform: str | None = None,
    has_email: bool | None = None,
    status: str | None = None,
    min_score: float | None = None,
    db: AsyncSession = Depends(get_db),
):
    """Export filtered creators as a styled Excel file (GET endpoint)."""
    filters = SearchFilters(
        q=q,
        category=category,
        language=language,
        creator_type=creator_type,
        broker=broker,
        audience_bucket=audience_bucket,
        platform=platform,
        has_email=has_email,
        status=status,
        min_score=min_score,
        page=1,
        page_size=5000,
    )

    service = CreatorService(db)
    result = await service.search_creators(filters)
    rows = [r.model_dump() if hasattr(r, "model_dump") else r for r in result.results]

    output = build_styled_excel(rows)
    filename = f"creators_export_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
